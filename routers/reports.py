from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from auth import current_user
from database import query


router = APIRouter(
    prefix="/api/export",
    tags=["Reports"]
)


def style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    ws.freeze_panes = "A2"


def auto_width(ws):
    for column in ws.columns:
        max_length = 0

        for cell in column:
            try:
                length = len(str(cell.value))
                if length > max_length:
                    max_length = length
            except Exception:
                pass

        width = min(max(max_length + 2, 12), 45)

        ws.column_dimensions[
            column[0].column_letter
        ].width = width


@router.get("/mine")
def export_mine(user=Depends(current_user)):

    user_id = user["UserId"]

    try:

        # ==================================================
        # 1. TẠO FILE EXCEL
        # ==================================================

        wb = Workbook()


        # ==================================================
        # 2. TỔNG HỢP
        # ==================================================

        ws = wb.active
        ws.title = "Tổng hợp"

        ws.append([
            "Thông tin",
            "Nội dung"
        ])

        ws.append([
            "Mã người dùng",
            user.get("UserId")
        ])

        ws.append([
            "Họ tên",
            user.get("FullName", "")
        ])

        ws.append([
            "Email",
            user.get("Email", "")
        ])

        ws.append([
            "Vai trò",
            user.get("Role", "")
        ])

        ws.append([])

        ws.append([
            "Hệ thống",
            "HealthBot - Hệ thống tư vấn sức khỏe"
        ])

        ws.append([
            "Nội dung",
            "Báo cáo dữ liệu sức khỏe cá nhân"
        ])

        style_header(ws)
        auto_width(ws)


        # ==================================================
        # 3. TRIỆU CHỨNG
        # ==================================================

        ws = wb.create_sheet("Triệu chứng")

        ws.append([
            "Mã triệu chứng",
            "Tên triệu chứng",
            "Mô tả",
            "Mức độ",
            "Thời gian tạo"
        ])

        rows = query(
            """
            SELECT
                SymptomId,
                SymptomName,
                Description,
                Severity,
                CreatedAt
            FROM Symptoms
            WHERE UserId = ?
            ORDER BY CreatedAt DESC
            """,
            (user_id,)
        )

        for row in rows:
            ws.append([
                row.get("SymptomId"),
                row.get("SymptomName"),
                row.get("Description"),
                row.get("Severity"),
                row.get("CreatedAt")
            ])

        style_header(ws)
        auto_width(ws)


        # ==================================================
        # 4. PHÂN TÍCH TRIỆU CHỨNG
        # ==================================================

        ws = wb.create_sheet("Phân tích triệu chứng")

        ws.append([
            "Mã kiểm tra",
            "Triệu chứng",
            "Mức độ",
            "Số ngày",
            "Kết quả phân tích",
            "Thời gian"
        ])

        rows = query(
            """
            SELECT
                CheckId,
                SymptomText,
                Severity,
                DaysSince,
                ResultJson,
                CreatedAt
            FROM SymptomChecks
            WHERE UserId = ?
            ORDER BY CreatedAt DESC
            """,
            (user_id,)
        )

        for row in rows:
            ws.append([
                row.get("CheckId"),
                row.get("SymptomText"),
                row.get("Severity"),
                row.get("DaysSince"),
                row.get("ResultJson"),
                row.get("CreatedAt")
            ])

        style_header(ws)
        auto_width(ws)


        # ==================================================
        # 5. THUỐC
        # ==================================================

        ws = wb.create_sheet("Thuốc")

        ws.append([
            "Mã thuốc",
            "Tên thuốc",
            "Liều dùng",
            "Số lần/ngày",
            "Ngày bắt đầu",
            "Ngày kết thúc",
            "Ghi chú",
            "Trạng thái"
        ])

        rows = query(
            """
            SELECT
                MedId,
                DrugName,
                Dosage,
                TimesPerDay,
                StartDate,
                EndDate,
                Note,
                IsActive
            FROM Medications
            WHERE UserId = ?
            ORDER BY MedId DESC
            """,
            (user_id,)
        )

        for row in rows:
            ws.append([
                row.get("MedId"),
                row.get("DrugName"),
                row.get("Dosage"),
                row.get("TimesPerDay"),
                row.get("StartDate"),
                row.get("EndDate"),
                row.get("Note"),
                row.get("IsActive")
            ])

        style_header(ws)
        auto_width(ws)


        # ==================================================
        # 6. GIỜ UỐNG THUỐC
        # ==================================================

        ws = wb.create_sheet("Lịch uống thuốc")

        ws.append([
            "Mã giờ",
            "Mã thuốc",
            "Tên thuốc",
            "Liều dùng",
            "Giờ uống"
        ])

        rows = query(
            """
            SELECT
                mt.TimeId,
                mt.MedId,
                m.DrugName,
                m.Dosage,
                mt.TakeTime
            FROM MedicationTimes mt
            INNER JOIN Medications m
                ON mt.MedId = m.MedId
            WHERE m.UserId = ?
            ORDER BY m.DrugName, mt.TakeTime
            """,
            (user_id,)
        )

        for row in rows:
            ws.append([
                row.get("TimeId"),
                row.get("MedId"),
                row.get("DrugName"),
                row.get("Dosage"),
                row.get("TakeTime")
            ])

        style_header(ws)
        auto_width(ws)


        # ==================================================
        # 7. LỊCH SỬ UỐNG THUỐC
        # ==================================================

        ws = wb.create_sheet("Lịch sử uống thuốc")

        ws.append([
            "Mã log",
            "Mã thuốc",
            "Tên thuốc",
            "Thời gian dự kiến",
            "Trạng thái",
            "Thời gian thực hiện",
            "Ảnh minh chứng"
        ])

        rows = query(
            """
            SELECT
                ml.LogId,
                ml.MedId,
                m.DrugName,
                ml.ScheduledAt,
                ml.Status,
                ml.ActionAt,
                ml.ProofImagePath
            FROM MedicationLogs ml
            INNER JOIN Medications m
                ON ml.MedId = m.MedId
            WHERE m.UserId = ?
            ORDER BY ml.ScheduledAt DESC
            """,
            (user_id,)
        )

        for row in rows:
            ws.append([
                row.get("LogId"),
                row.get("MedId"),
                row.get("DrugName"),
                row.get("ScheduledAt"),
                row.get("Status"),
                row.get("ActionAt"),
                row.get("ProofImagePath")
            ])

        style_header(ws)
        auto_width(ws)


        # ==================================================
        # 8. LỊCH TÁI KHÁM
        # ==================================================

        ws = wb.create_sheet("Tái khám")

        ws.append([
            "Mã lịch",
            "Bác sĩ",
            "Địa điểm",
            "Thời gian",
            "Nội dung",
            "Ghi chú",
            "Trạng thái",
            "Nhắc trước",
            "Kênh",
            "Xác nhận lúc",
            "Số lần nhắc",
            "Ngày tạo"
        ])

        rows = query(
            """
            SELECT
                ApptId,
                DoctorName,
                Location,
                ApptTime,
                Content,
                Note,
                Status,
                RemindBefore,
                Channel,
                ConfirmedAt,
                ReminderCount,
                CreatedAt
            FROM Appointments
            WHERE UserId = ?
            ORDER BY ApptTime DESC
            """,
            (user_id,)
        )

        for row in rows:
            ws.append([
                row.get("ApptId"),
                row.get("DoctorName"),
                row.get("Location"),
                row.get("ApptTime"),
                row.get("Content"),
                row.get("Note"),
                row.get("Status"),
                row.get("RemindBefore"),
                row.get("Channel"),
                row.get("ConfirmedAt"),
                row.get("ReminderCount"),
                row.get("CreatedAt")
            ])

        style_header(ws)
        auto_width(ws)


        # ==================================================
        # 9. THÔNG BÁO
        # ==================================================

        ws = wb.create_sheet("Thông báo")

        ws.append([
            "Mã thông báo",
            "Tiêu đề",
            "Nội dung",
            "Loại",
            "Thời gian gửi",
            "Trạng thái"
        ])

        rows = query(
            """
            SELECT
                NotificationId,
                Title,
                Message,
                NotificationType,
                SentAt,
                Status
            FROM Notifications
            WHERE UserId = ?
            ORDER BY SentAt DESC
            """,
            (user_id,)
        )

        for row in rows:
            ws.append([
                row.get("NotificationId"),
                row.get("Title"),
                row.get("Message"),
                row.get("NotificationType"),
                row.get("SentAt"),
                row.get("Status")
            ])

        style_header(ws)
        auto_width(ws)


        # ==================================================
        # 10. LỊCH SỬ CHATBOT
        # ==================================================

        ws = wb.create_sheet("Chatbot")

        ws.append([
            "Mã chat",
            "Tin nhắn người dùng",
            "Phản hồi AI",
            "Thời gian"
        ])

        rows = query(
            """
            SELECT
                ID,
                Message,
                Response,
                CreatedAt
            FROM ChatHistory
            WHERE UserID = ?
            ORDER BY CreatedAt DESC
            """,
            (user_id,)
        )

        for row in rows:
            ws.append([
                row.get("ID"),
                row.get("Message"),
                row.get("Response"),
                row.get("CreatedAt")
            ])

        style_header(ws)
        auto_width(ws)


        # ==================================================
        # 11. TẠO FILE
        # ==================================================

        output = BytesIO()

        wb.save(output)

        output.seek(0)


        # ==================================================
        # 12. TRẢ FILE CHO TRÌNH DUYỆT
        # ==================================================

        return StreamingResponse(
            output,
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition":
                    'attachment; filename="HealthBot_BaoCao.xlsx"'
            }
        )


    except Exception as e:

        print("\n================================")
        print("LỖI XUẤT BÁO CÁO")
        print("================================")
        print(type(e).__name__)
        print(str(e))
        print("================================\n")

        raise HTTPException(
            status_code=500,
            detail=f"Lỗi xuất báo cáo: {str(e)}"
        )